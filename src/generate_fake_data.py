# -*- coding: utf-8 -*-
"""
演示数据生成器
=====================================================================
生成一份结构与真实业务模板完全一致、金额全部虚构的
《预算使用明细表.xlsx》（含 汇总 / 授予额度 / 差旅费 /
费用明细-费用明细 / 专业费用 五个 sheet），供一键核查脚本演示。

刻意埋入三类"考点"，用于演示脚本的六项检查能力：
  1. 真实超支 3 处（业务问题，非表错误）：
     招待费-归口报账@重庆、市场推广费@云南、专业费用@四川
  2. 未归类长科目名 1 处（尾差分解演示）：
     费用明细中「设备租赁费-IT归口办公电脑租赁」@贵州
  3. 汇总表数据与底表完全勾稽（逐格重算应 0 差异）

用法：
  python generate_fake_data.py            # 默认生成 9 月
  python generate_fake_data.py --month 10 # 生成 10 月目录
"""
import os
import random
import argparse
import openpyxl
from openpyxl.utils import get_column_letter

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(REPO_ROOT, "data", "预算使用明细表")

REGIONS = ["重庆", "云南", "贵州", "四川", "西藏", "运营部"]

# ---------------- 汇总表科目行（顺序须与标准科目清单一致） ----------------
# (科目名, 来源)  来源: fy=费用明细表 / zy_parent=专业费用父行 /
#                   zy_child=专业费用子行(按I列预算项目) / zy_it=专业费用-IT归口 / cl=差旅费表
SUBJECT_ROWS = [
    ("差旅费",                        "cl"),
    ("招待费-归口报账",              "fy"),
    ("招待费-员工报账",              "fy"),
    ("市场推广费",                   "fy"),
    ("专业费用",                     "zy_parent"),
    ("专业费用-劳务费-专家咨询费",   "zy_child"),
    ("专业费用-咨询费",              "zy_child"),
    ("邮寄运输费",                   "fy"),
    ("办公类",                       "fy"),
    ("低值易耗品消耗",               "fy"),
    ("设备维修费",                   "fy"),
    ("其他费用",                     "fy"),
    ("专业费用-IT归口",              "zy_it"),
]

# ---------------- 授予额度（万元）----------------
GRANT = {
    "差旅费":             {"重庆": 12, "云南": 14, "贵州": 10, "四川": 16, "西藏": 8,  "运营部": 20},
    "招待费-归口报账":   {"重庆": 8,  "云南": 6,  "贵州": 5,  "四川": 7,  "西藏": 4,  "运营部": 10},
    "招待费-员工报账":   {"重庆": 6,  "云南": 5,  "贵州": 4,  "四川": 6,  "西藏": 3,  "运营部": 8},
    "市场推广费":        {"重庆": 12, "云南": 15, "贵州": 8,  "四川": 18, "西藏": 6,  "运营部": 20},
    "邮寄运输费":        {"重庆": 5,  "云南": 4,  "贵州": 3,  "四川": 6,  "西藏": 2,  "运营部": 6},
    "办公类":            {"重庆": 6,  "云南": 5,  "贵州": 4,  "四川": 6,  "西藏": 3,  "运营部": 7},
    "低值易耗品消耗":    {"重庆": 3,  "云南": 2,  "贵州": 2,  "四川": 3,  "西藏": 1,  "运营部": 4},
    "设备维修费":        {"重庆": 4,  "云南": 3,  "贵州": 2,  "四川": 5,  "西藏": 2,  "运营部": 5},
    "其他费用":          {"重庆": 2,  "云南": 2,  "贵州": 1,  "四川": 2,  "西藏": 1,  "运营部": 3},
    "专业费用":          {"重庆": 25, "云南": 20, "贵州": 18, "四川": 40, "西藏": 15, "运营部": 35},
    "专业费用-IT归口":   {"重庆": 8,  "云南": 6,  "贵州": 5,  "四川": 8,  "西藏": 4,  "运营部": 9},
}

# ---------------- 已使用目标值（万元）----------------
USED = {
    "差旅费":             {"重庆": 8.5,  "云南": 10.2, "贵州": 6.8, "四川": 11.0, "西藏": 5.5, "运营部": 13.0},
    "招待费-归口报账":   {"重庆": 10.2, "云南": 4.2,  "贵州": 3.1, "四川": 5.0,  "西藏": 2.8, "运营部": 6.5},
    "招待费-员工报账":   {"重庆": 4.1,  "云南": 3.5,  "贵州": 2.8, "四川": 4.2,  "西藏": 2.0, "运营部": 5.5},
    "市场推广费":        {"重庆": 8.5,  "云南": 19.5, "贵州": 5.5, "四川": 12.0, "西藏": 4.0, "运营部": 13.5},
    "邮寄运输费":        {"重庆": 3.2,  "云南": 2.6,  "贵州": 1.9, "四川": 3.8,  "西藏": 1.4, "运营部": 4.0},
    "办公类":            {"重庆": 4.2,  "云南": 3.5,  "贵州": 2.8, "四川": 4.0,  "西藏": 2.1, "运营部": 4.8},
    "低值易耗品消耗":    {"重庆": 1.8,  "云南": 1.2,  "贵州": 1.0, "四川": 1.9,  "西藏": 0.6, "运营部": 2.4},
    "设备维修费":        {"重庆": 2.5,  "云南": 1.8,  "贵州": 1.2, "四川": 3.1,  "西藏": 1.3, "运营部": 3.2},
    "其他费用":          {"重庆": 1.2,  "云南": 1.0,  "贵州": 0.6, "四川": 1.3,  "西藏": 0.7, "运营部": 1.8},
    "专业费用-劳务费-专家咨询费": {"重庆": 12.0, "云南": 9.0, "贵州": 8.0, "四川": 30.0, "西藏": 7.0, "运营部": 16.0},
    "专业费用-咨询费":            {"重庆": 6.0,  "云南": 5.0, "贵州": 4.0, "四川": 14.0, "西藏": 4.0, "运营部": 9.0},
    "专业费用-IT归口":            {"重庆": 5.2,  "云南": 3.8, "贵州": 3.1, "四川": 5.0,  "西藏": 2.6, "运营部": 6.0},
}

# 汇总表区域 -> "已使用"列
REGION_USE_COL = {"运营部": "F", "重庆": "I", "云南": "L",
                  "贵州": "O", "四川": "R", "西藏": "U"}


def split_amount(total_wan, n, rng):
    """把 total_wan 万元拆成 n 笔元金额（整数元，合计精确）"""
    total_yuan = int(round(total_wan * 10000))
    if n <= 1 or total_yuan <= 0:
        return [total_yuan]
    cuts = sorted(rng.randint(1, total_yuan - 1) for _ in range(n - 1))
    parts, prev = [], 0
    for c in cuts:
        parts.append(c - prev)
        prev = c
    parts.append(total_yuan - prev)
    return parts


def build_details(rng):
    """生成三张底表的明细行。返回 (差旅行, 费用明细行, 专业费用行)"""
    cl_rows, fy_rows, zy_rows = [], [], []

    # 差旅费表：U=不含税金额(元)，AF=区域
    for reg in REGIONS:
        for amt in split_amount(USED["差旅费"][reg], rng.randint(3, 6), rng):
            cl_rows.append({"amount": amt, "region": reg})

    # 费用明细表：G=预算科目，T=发生金额(元)，V=区域
    fy_subjects = ["招待费-归口报账", "招待费-员工报账", "市场推广费", "邮寄运输费",
                   "办公类", "低值易耗品消耗", "设备维修费", "其他费用"]
    for subj in fy_subjects:
        for reg in REGIONS:
            for amt in split_amount(USED[subj][reg], rng.randint(2, 5), rng):
                fy_rows.append({"subject": subj, "amount": amt, "region": reg})
    # 埋点：未归类长科目名（汇总表无此科目行 → 尾差分解演示）
    fy_rows.append({"subject": "设备租赁费-IT归口办公电脑租赁",
                    "amount": 46000, "region": "贵州"})

    # 专业费用表：G=预算科目，I=预算项目，V=发生金额(元)，X=区域
    for child in ["专业费用-劳务费-专家咨询费", "专业费用-咨询费"]:
        for reg in REGIONS:
            for amt in split_amount(USED[child][reg], rng.randint(2, 4), rng):
                zy_rows.append({"subject": "专业费用", "project": child,
                                "amount": amt, "region": reg})
    for reg in REGIONS:
        for amt in split_amount(USED["专业费用-IT归口"][reg], rng.randint(1, 3), rng):
            zy_rows.append({"subject": "专业费用-IT归口", "project": "IT运维服务",
                            "amount": amt, "region": reg})
    return cl_rows, fy_rows, zy_rows


def used_map():
    """从 USED 推导各科目行的区域已用值（万元），与底表口径一致"""
    m = {}
    for name, src in SUBJECT_ROWS:
        if src == "zy_parent":
            m[name] = {r: USED["专业费用-劳务费-专家咨询费"][r] + USED["专业费用-咨询费"][r]
                       for r in REGIONS}
        else:
            m[name] = dict(USED[name])
    return m


def write_sheet_headers(ws, headers):
    """headers: {列字母: 表头文字}"""
    for letter, text in headers.items():
        ws.cell(row=1, column=openpyxl.utils.column_index_from_string(letter), value=text)


def build_workbook(rng):
    wb = openpyxl.Workbook()
    um = used_map()

    # ---- 汇总 ----
    ws = wb.active
    ws.title = "汇总"
    ws.cell(row=1, column=1, value="科目")
    ws.cell(row=2, column=2, value="授予")
    ws.cell(row=2, column=3, value="已使用")
    ws.cell(row=2, column=4, value="剩余")
    for reg, L in REGION_USE_COL.items():
        c = openpyxl.utils.column_index_from_string(L)
        ws.cell(row=1, column=c, value=reg)
        ws.cell(row=2, column=c, value="已使用")

    r = 3
    for name, src in SUBJECT_ROWS:
        ws.cell(row=r, column=1, value=name)
        region_used = um[name]
        if src != "zy_child":                      # 展开子行不写授予/剩余（防重复计数）
            grant_total = sum(GRANT[name].values())
            used_total = sum(region_used.values())
            ws.cell(row=r, column=2, value=grant_total)
            ws.cell(row=r, column=3, value=round(used_total, 4))
            ws.cell(row=r, column=4, value=round(grant_total - used_total, 4))
        for reg, L in REGION_USE_COL.items():
            c = openpyxl.utils.column_index_from_string(L)
            ws.cell(row=r, column=c, value=round(region_used[reg], 4))
        r += 1
    # 总计行（不含子行，防重复）
    ws.cell(row=r, column=1, value="总计")
    ws.cell(row=r, column=2, value=sum(sum(GRANT[n].values()) for n, s in SUBJECT_ROWS if s != "zy_child"))
    used_all = sum(sum(um[n].values()) for n, s in SUBJECT_ROWS if s != "zy_child")
    ws.cell(row=r, column=3, value=round(used_all, 4))
    ws.cell(row=r, column=4, value=round(ws.cell(row=r, column=2).value - used_all, 4))
    for reg, L in REGION_USE_COL.items():
        c = openpyxl.utils.column_index_from_string(L)
        v = sum(um[n][reg] for n, s in SUBJECT_ROWS if s != "zy_child")
        ws.cell(row=r, column=c, value=round(v, 4))

    # ---- 授予额度 ----
    ws = wb.create_sheet("授予额度")
    ws.cell(row=3, column=1, value="科目")
    for i, reg in enumerate(REGIONS):
        ws.cell(row=3, column=2 + i, value=reg)
    r = 4
    for name in GRANT:
        ws.cell(row=r, column=1, value=name)
        for i, reg in enumerate(REGIONS):
            ws.cell(row=r, column=2 + i, value=GRANT[name][reg])
        r += 1

    # ---- 三张底表 ----
    cl_rows, fy_rows, zy_rows = build_details(rng)

    ws = wb.create_sheet("差旅费")
    write_sheet_headers(ws, {"A": "序号", "B": "出差人", "C": "部门", "D": "出差事由",
                             "U": "不含税金额CNY", "AF": "区域"})
    for i, row in enumerate(cl_rows, start=2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=21, value=row["amount"])
        ws.cell(row=i, column=32, value=row["region"])

    ws = wb.create_sheet("费用明细-费用明细")
    write_sheet_headers(ws, {"A": "序号", "G": "预算科目", "T": "发生金额",
                             "V": "区域", "W": "备注"})
    for i, row in enumerate(fy_rows, start=2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=7, value=row["subject"])
        ws.cell(row=i, column=20, value=row["amount"])
        ws.cell(row=i, column=22, value=row["region"])

    ws = wb.create_sheet("专业费用")
    write_sheet_headers(ws, {"A": "序号", "G": "预算科目", "I": "预算项目",
                             "V": "发生金额", "X": "区域"})
    for i, row in enumerate(zy_rows, start=2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=7, value=row["subject"])
        ws.cell(row=i, column=9, value=row["project"])
        ws.cell(row=i, column=22, value=row["amount"])
        ws.cell(row=i, column=24, value=row["region"])

    return wb


def main():
    ap = argparse.ArgumentParser(description="生成演示用《预算使用明细表.xlsx》")
    ap.add_argument("--month", type=int, default=9, help="月份数字（默认 9）")
    args = ap.parse_args()

    rng = random.Random(2026)          # 固定种子：演示可复现
    month_dir = os.path.join(DATA_DIR, f"{args.month}月")
    os.makedirs(month_dir, exist_ok=True)
    path = os.path.join(month_dir, "预算使用明细表.xlsx")

    wb = build_workbook(rng)
    wb.save(path)
    print(f"[OK] 演示数据已生成: {path}")
    print("     内含 3 处真实超支 + 1 处未归类科目（尾差演示），其余全部勾稽。")


if __name__ == "__main__":
    main()
